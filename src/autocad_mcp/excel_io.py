"""Leitura/escrita de planilhas e mapeamento de colunas — 100% puro (sem AutoCAD), testável isoladamente."""

from __future__ import annotations

import os
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from autocad_mcp.config import Config
from autocad_mcp.errors import ExcelIntegrationError, InvalidParameterError
from autocad_mcp.geometry import format_cell, parse_number
from autocad_mcp.models import RowSpec

FIXED_HEADERS = ["HANDLE", "BLOCO", "CAMADA", "X", "Y", "Z", "ROTACAO", "ESCALA_X", "ESCALA_Y", "ESCALA_Z"]

_ALIASES = {
    "handle": "handle",
    "x": "x", "y": "y", "z": "z",
    "rotation": "rotation", "rotacao": "rotation", "rot": "rotation",
    "scale": "scale", "escala": "scale",
    "scale_x": "scale_x", "escala_x": "scale_x", "scale_y": "scale_y", "escala_y": "scale_y",
    "scale_z": "scale_z", "escala_z": "scale_z",
    "layer": "layer", "camada": "layer",
    "block": "block", "bloco": "block", "block_name": "block", "nome_bloco": "block",
}
_TARGETS = set(_ALIASES.values()) | {"ignore"}


def norm_header(h: Any) -> str:
    text = unicodedata.normalize("NFKD", str(h).strip()).encode("ascii", "ignore").decode()
    return text.casefold().replace(" ", "_")


@dataclass
class Table:
    sheet: str
    headers: list[str]                    # cabeçalhos originais (strip)
    rows: list[tuple[int, dict[str, Any]]]  # (nº da linha no Excel, {cabeçalho: valor})
    warnings: list[str] = field(default_factory=list)


def check_path(path: str, cfg: Config, *, write: bool) -> Path:
    if not path or not str(path).strip():
        raise ExcelIntegrationError("Informe o caminho da planilha.")
    p = Path(os.path.expandvars(os.path.expanduser(str(path)))).resolve()
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelIntegrationError(f"Extensão '{p.suffix}' não suportada: use .xlsx (ou .xlsm para leitura). .xls/.csv não são lidos.")
    if cfg.excel_allowed_dirs and not any(p.is_relative_to(d.resolve()) for d in cfg.excel_allowed_dirs):
        raise ExcelIntegrationError(f"Caminho fora das pastas permitidas em [excel].allowed_dirs: {[str(d) for d in cfg.excel_allowed_dirs]}")
    if write:
        if p.suffix.lower() != ".xlsx":
            raise ExcelIntegrationError("Só é possível gravar .xlsx.")
        p.parent.mkdir(parents=True, exist_ok=True)
    elif not p.is_file():
        raise ExcelIntegrationError(f"Arquivo não encontrado: {p}")
    return p


def read_table(path: Path, sheet: str | None, header_row: int, max_rows: int) -> Table:
    if header_row < 1:
        raise InvalidParameterError("header_row começa em 1.")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - arquivo corrompido/aberto/bloqueado
        raise ExcelIntegrationError(f"Não consegui abrir a planilha ({type(exc).__name__}: {exc}). Ela está corrompida ou protegida?") from exc
    try:
        names = wb.sheetnames
        if sheet:
            match = next((n for n in names if n.casefold() == sheet.strip().casefold()), None)
            if match is None:
                raise ExcelIntegrationError(f"Aba '{sheet}' não existe. Abas: {names}")
        else:
            match = names[0]
        ws = wb[match]
        headers: list[str] = []
        col_idx: list[int] = []
        rows: list[tuple[int, dict[str, Any]]] = []
        warnings: list[str] = []
        for rnum, values in enumerate(ws.iter_rows(min_row=header_row, values_only=True), start=header_row):
            if rnum == header_row:
                for i, v in enumerate(values):
                    if v is not None and str(v).strip():
                        headers.append(str(v).strip())
                        col_idx.append(i)
                continue
            rec = {h: (values[i] if i < len(values) else None) for h, i in zip(headers, col_idx)}
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in rec.values()):
                continue
            rows.append((rnum, rec))
            if len(rows) > max_rows:
                raise ExcelIntegrationError(f"A planilha tem mais de {max_rows} linhas de dados (limite `max_batch_rows`). Divida o arquivo.")
        if not headers:
            raise ExcelIntegrationError(f"Nenhum cabeçalho encontrado na linha {header_row} da aba '{match}'.")
        seen: dict[str, str] = {}
        for h in headers:
            k = norm_header(h)
            if k in seen:
                raise ExcelIntegrationError(f"Cabeçalhos duplicados na planilha: '{seen[k]}' e '{h}'.")
            seen[k] = h
        if not rows:
            warnings.append("A planilha não tem linhas de dados abaixo do cabeçalho.")
        return Table(match, headers, rows, warnings)
    finally:
        wb.close()


def write_table(path: Path, sheet: str, headers: list[str], rows: list[list[Any]], info: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31] or "Blocos"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(format_cell(r[i - 1])) for r in rows[:200] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)
    meta = wb.create_sheet("INFO")
    for k, v in info.items():
        meta.append([k, format_cell(v) if not isinstance(v, (int, float)) else v])
    meta.column_dimensions["A"].width = 22
    meta.column_dimensions["B"].width = 60
    wb.save(path)


# --- mapeamento de colunas -------------------------------------------------------------------
def resolve_mapping(headers: list[str], mapping: dict[str, str] | None) -> dict[str, str]:
    """{cabeçalho: alvo}. alvo ∈ campos fixos | 'attr:TAG' | 'ignore'. Sem `mapping`: detecção automática."""
    out: dict[str, str] = {}
    if mapping:
        by_norm = {norm_header(h): h for h in headers}
        for col, target in mapping.items():
            h = by_norm.get(norm_header(col))
            if h is None:
                raise InvalidParameterError(f"Coluna '{col}' do mapping não existe na planilha. Colunas: {headers}")
            t = str(target).strip()
            if t.casefold().startswith("attr:"):
                tag = t[5:].strip().upper()
                if not tag:
                    raise InvalidParameterError(f"Alvo '{target}' sem TAG após 'attr:'.")
                out[h] = "attr:" + tag
            elif norm_header(t) in _TARGETS:
                out[h] = _ALIASES.get(norm_header(t), "ignore")
            else:
                raise InvalidParameterError(
                    f"Alvo de mapping inválido: '{target}'. Use x,y,z,rotation,scale,scale_x/y/z,layer,block,handle,ignore ou attr:TAG."
                )
        return out
    for h in headers:
        k = norm_header(h)
        if k.startswith("attr:"):
            out[h] = "attr:" + h.strip()[5:].strip().upper()
        else:
            out[h] = _ALIASES.get(k) or "attr:" + h.strip().upper()
    return out


def _cell_number(v: Any, what: str) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return parse_number(v, name=what)


def build_rows(table: Table, mapping: dict[str, str], default_block: str | None) -> tuple[list[RowSpec], list[dict[str, Any]]]:
    """Linhas normalizadas + erros por linha ({row, action:'error', message}). Célula em branco = ignorar."""
    specs: list[RowSpec] = []
    errors: list[dict[str, Any]] = []
    for rnum, rec in table.rows:
        try:
            data: dict[str, Any] = {"row": rnum, "attributes": {}}
            scale = [1.0, 1.0, 1.0]
            for h, target in mapping.items():
                v = rec.get(h)
                blank = v is None or (isinstance(v, str) and not v.strip())
                if target == "ignore" or blank:
                    continue
                if target.startswith("attr:"):
                    data["attributes"][target[5:]] = format_cell(v)
                elif target in ("x", "y", "z", "rotation"):
                    data[target] = _cell_number(v, f"{h} (linha {rnum})")
                elif target == "scale":
                    scale = [_cell_number(v, f"{h} (linha {rnum})")] * 3
                elif target in ("scale_x", "scale_y", "scale_z"):
                    scale["xyz".index(target[-1])] = _cell_number(v, f"{h} (linha {rnum})")
                elif target in ("layer", "block", "handle"):
                    data[target] = format_cell(v)
            data["scale"] = tuple(scale)
            if not data.get("block") and default_block:
                data["block"] = default_block
            specs.append(RowSpec(**data))
        except (InvalidParameterError, ValidationError) as exc:
            errors.append({"row": rnum, "action": "error", "message": str(exc).splitlines()[0]})
    return specs, errors
