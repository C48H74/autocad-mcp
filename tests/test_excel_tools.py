"""Etapa 4: import/export/sync com AutoCAD falso — inclui o critério de aceitação dos 100 blocos + 1 Ctrl+Z."""

import pytest
from openpyxl import Workbook, load_workbook

from tests.test_excel import make_xlsx


@pytest.fixture
def edoc(doc):
    doc.Layers.Add("SUPORTES")
    doc.Blocks.define("SUP-GUIDE", {"TAG": "", "TIPO": "GUIA", "LINHA": ""})
    doc.Blocks.define("SUP-SHOE", {"TAG": "", "TIPO": "SHOE"})
    doc.journal.groups.clear()
    return doc


def table_100(tmp_path):
    rows = [["TAG", "TIPO", "X", "Y", "Z", "LINHA", "CAMADA"]]
    for i in range(100):
        rows.append([f"PS-{i:03d}", "GUIA" if i % 2 else "SHOE", i * 500.0, 1000 + i, 4200, f"10-P-{1000 + i}", "SUPORTES"])
    return make_xlsx(tmp_path / "suportes.xlsx", rows)


def test_import_100_blocks_dry_run_then_real_then_single_ctrl_z(call, edoc, tmp_path):
    p = table_100(tmp_path)
    dry = call("import_blocks_from_excel", path=str(p), block_name="SUP-GUIDE", key_tag="TAG", dry_run=True)
    assert dry["ok"] and dry["data"]["to_insert"] == 100 and dry["data"]["errors"] == 0 and edoc.model.Count == 0
    assert dry["data"]["units"] == "mm" and edoc.journal.groups == []  # dry_run não abre marca de undo

    r = call("import_blocks_from_excel", path=str(p), block_name="SUP-GUIDE", key_tag="TAG")
    assert r["ok"] and r["data"]["inserted"] == 100 and r["data"]["errors"] == 0 and r["data"]["applied"]
    assert edoc.model.Count == 100
    blocks = call("read_block_attributes", limit=500)["data"]["blocks"]
    b = next(x for x in blocks if x["attributes"]["TAG"] == "PS-042")
    assert b["position"] == [21000.0, 1042.0, 4200.0] and b["layer"] == "SUPORTES" and b["attributes"]["LINHA"] == "10-P-1042"
    assert len(edoc.journal.groups) == 1  # 100 inserções + atributos = UM grupo de undo
    assert edoc.ctrl_z() and edoc.model.Count == 0  # critério de aceitação: um único Ctrl+Z desfaz tudo
    assert not edoc.ctrl_z()


def test_reimport_same_file_is_idempotent_thanks_to_key_tag(call, edoc, tmp_path):
    p = table_100(tmp_path)
    call("import_blocks_from_excel", path=str(p), block_name="SUP-GUIDE", key_tag="TAG")
    again = call("import_blocks_from_excel", path=str(p), block_name="SUP-GUIDE", key_tag="TAG")
    assert again["ok"] and again["data"]["unchanged"] == 100 and again["data"]["to_insert"] == 0 and edoc.model.Count == 100


def test_update_existing_requires_confirm_and_reports_position_divergence(call, edoc, tmp_path):
    p = make_xlsx(tmp_path / "a.xlsx", [["TAG", "TIPO", "X", "Y"], ["PS-1", "GUIA", 0, 0]])
    call("import_blocks_from_excel", path=str(p), block_name="SUP-GUIDE", key_tag="TAG")
    p2 = make_xlsx(tmp_path / "b.xlsx", [["TAG", "TIPO", "X", "Y"], ["PS-1", "ANCORA", 10, 10]])
    blocked = call("import_blocks_from_excel", path=str(p2), block_name="SUP-GUIDE", key_tag="TAG")
    assert not blocked["ok"] and blocked["error"]["code"] == "confirmation_required"
    assert blocked["error"]["details"]["to_update"] == 1
    ok = call("import_blocks_from_excel", path=str(p2), block_name="SUP-GUIDE", key_tag="TAG", confirm=True)
    row = ok["data"]["results"][0]
    assert ok["ok"] and ok["data"]["updated"] == 1 and row["changes"]["TIPO"] == {"from": "GUIA", "to": "ANCORA"}
    assert "update_position" in row["warning"]
    cur = call("read_block_attributes")["data"]["blocks"][0]
    assert cur["attributes"]["TIPO"] == "ANCORA" and cur["position"] == [0.0, 0.0, 0.0]  # posição preservada
    moved = call("import_blocks_from_excel", path=str(p2), block_name="SUP-GUIDE", key_tag="TAG", confirm=True, update_position=True)
    assert moved["data"]["updated"] == 1
    assert call("read_block_attributes")["data"]["blocks"][0]["position"] == [10.0, 10.0, 0.0]


def test_row_errors_are_reported_per_row_and_do_not_abort_batch(call, edoc, tmp_path):
    p = make_xlsx(tmp_path / "e.xlsx", [
        ["BLOCO", "TAG", "X", "Y", "CAMADA"],
        ["SUP-GUIDE", "OK-1", 1, 1, "SUPORTES"],
        ["NAO-EXISTE", "E-2", 1, 1, "SUPORTES"],
        ["SUP-GUIDE", "E-3", "abc", 1, "SUPORTES"],
        ["SUP-GUIDE", "E-4", None, None, "SUPORTES"],
        ["SUP-GUIDE", "E-5", 1, 1, "CAMADA-FANTASMA"],
        ["SUP-GUIDE", "OK-6", "1.234,5", "2,5", "SUPORTES"],
        ["SUP-GUIDE", "OK-1", 5, 5, "SUPORTES"],
    ])
    r = call("import_blocks_from_excel", path=str(p), key_tag="TAG")
    d = r["data"]
    assert r["ok"] and d["inserted"] == 2 and d["errors"] == 5 and edoc.model.Count == 2
    msgs = {e["row"]: e["message"] for e in d["error_rows"]}
    assert "não existe" in msgs[3] and "abc" in msgs[4] and "coordenadas" in msgs[5] and "não existe" in msgs[6] and "repetida" in msgs[8]
    six = next(x for x in call("read_block_attributes")["data"]["blocks"] if x["attributes"]["TAG"] == "OK-6")
    assert six["position"][:2] == [1234.5, 2.5]


def test_custom_mapping_default_layer_and_missing_layer_creation(call, edoc, tmp_path):
    p = make_xlsx(tmp_path / "m.xlsx", [["Suporte", "Este", "Norte", "Cota", "Giro"], ["S-9", 10, 20, 30, 90]])
    m = {"Suporte": "attr:TAG", "Este": "x", "Norte": "y", "Cota": "z", "Giro": "rotation"}
    fail = call("import_blocks_from_excel", path=str(p), mapping=m, block_name="SUP-SHOE", layer="NOVA")
    assert not fail["ok"] and fail["error"]["code"] == "invalid_parameter"
    ok = call("import_blocks_from_excel", path=str(p), mapping=m, block_name="SUP-SHOE", layer="NOVA", create_missing_layer=True)
    b = call("read_block_attributes")["data"]["blocks"][0]
    assert ok["ok"] and b["position"] == [10.0, 20.0, 30.0] and b["rotation_deg"] == 90.0 and b["layer"] == "NOVA"
    assert edoc.ctrl_z() and "NOVA" not in edoc.Layers.names() and edoc.model.Count == 0


def test_export_then_reimport_roundtrip(call, edoc, tmp_path):
    call("import_blocks_from_excel", path=str(table_100(tmp_path)), block_name="SUP-GUIDE", key_tag="TAG")
    out = tmp_path / "saida" / "export.xlsx"
    r = call("export_blocks_to_excel", block_name="SUP-*", path=str(out), layer="SUPORTES")
    assert r["ok"] and r["data"]["rows"] == 100 and r["data"]["columns"][:10] == ["HANDLE", "BLOCO", "CAMADA", "X", "Y", "Z", "ROTACAO", "ESCALA_X", "ESCALA_Y", "ESCALA_Z"]
    assert r["data"]["units"] == "mm"
    info = {row[0].value: row[1].value for row in load_workbook(out)["INFO"].iter_rows()}
    assert info["unidade"] == "mm" and info["desenho"] == "SUPORTES-01.dwg"
    again = call("import_blocks_from_excel", path=str(out), key_tag="TAG", dry_run=True)
    assert again["ok"] and again["data"]["unchanged"] == 100 and again["data"]["to_insert"] == 0 and again["data"]["errors"] == 0


def test_export_overwrite_needs_confirm_and_colliding_tag_names_are_prefixed(call, edoc, tmp_path):
    edoc.Blocks.define("ODD", {"X": "1", "TAG": ""})
    edoc.model.InsertBlock([0, 0, 0], "ODD", 1, 1, 1, 0)
    out = tmp_path / "e.xlsx"
    ok = call("export_blocks_to_excel", block_name="ODD", path=str(out))
    assert ok["ok"] and "ATTR:X" in ok["data"]["columns"] and "TAG" in ok["data"]["columns"]
    blocked = call("export_blocks_to_excel", block_name="ODD", path=str(out))
    assert not blocked["ok"] and blocked["error"]["code"] == "confirmation_required"
    assert call("export_blocks_to_excel", block_name="ODD", path=str(out), confirm=True)["ok"]
    empty = call("export_blocks_to_excel", block_name="NADA", path=str(tmp_path / "v.xlsx"))
    assert empty["ok"] and empty["data"]["rows"] == 0 and any("Nenhum bloco" in w for w in empty["warnings"])
    assert not call("export_blocks_to_excel", block_name="ODD", path=str(tmp_path / "x.csv"))["ok"]


def test_sync_attributes_dry_run_confirm_undo_and_report(call, edoc, tmp_path):
    call("import_blocks_from_excel", path=str(table_100(tmp_path)), block_name="SUP-GUIDE", key_tag="TAG")
    edoc.model.InsertBlock([0, 0, 0], "SUP-GUIDE", 1, 1, 1, 0).GetAttributes()[0].TextString = "SO-NO-DESENHO"
    edoc.model.InsertBlock([0, 0, 0], "SUP-GUIDE", 1, 1, 1, 0).GetAttributes()[0].TextString = "PS-005"  # duplicata da chave
    edoc.journal.groups.clear()
    wb = Workbook()
    ws = wb.active
    ws.append(["TAG", "TIPO", "LINHA", "PESO"])
    ws.append(["PS-000", "ANCORA", None, 12])       # muda TIPO; LINHA em branco mantém; PESO não é tag do bloco
    ws.append(["PS-001", "GUIA", "10-P-1001", None])  # nada muda
    ws.append(["PS-404", "X", "Y", None])           # não encontrado
    ws.append(["PS-005", "X", "Y", None])           # ambíguo (2 no desenho)
    ws.append([None, "X", "Y", None])               # chave vazia
    ws.append(["PS-000", "X", "Y", None])           # chave repetida na planilha
    p = tmp_path / "sync.xlsx"
    wb.save(p)

    dry = call("sync_attributes_from_excel", path=str(p), key_tag="TAG", dry_run=True)
    d = dry["data"]
    assert dry["ok"] and (d["changed"], d["unchanged"], d["not_found"], d["ambiguous"], d["error"]) == (1, 1, 1, 1, 2)
    assert d["results"][0]["changes"] == {"TIPO": {"from": "SHOE", "to": "ANCORA"}}
    assert d["in_drawing_not_in_sheet"] == 98  # 100 chaves PS-* + "so-no-desenho" − {PS-000, PS-001, PS-005} presentes na planilha
    assert any("PESO" in w for w in dry["warnings"])
    assert edoc.journal.groups == []

    blocked = call("sync_attributes_from_excel", path=str(p), key_tag="TAG")
    assert not blocked["ok"] and blocked["error"]["code"] == "confirmation_required"
    done = call("sync_attributes_from_excel", path=str(p), key_tag="TAG", confirm=True)
    assert done["ok"] and done["data"]["applied"] and len(edoc.journal.groups) == 1
    ps0 = next(x for x in call("read_block_attributes", limit=500)["data"]["blocks"] if x["attributes"]["TAG"] == "PS-000")
    assert ps0["attributes"]["TIPO"] == "ANCORA" and ps0["attributes"]["LINHA"] == "10-P-1000"
    assert edoc.ctrl_z()
    ps0 = next(x for x in call("read_block_attributes", limit=500)["data"]["blocks"] if x["attributes"]["TAG"] == "PS-000")
    assert ps0["attributes"]["TIPO"] == "SHOE"


def test_sync_requires_key_column_and_valid_file(call, edoc, tmp_path):
    p = make_xlsx(tmp_path / "n.xlsx", [["NOME", "TIPO"], ["a", "b"]])
    r = call("sync_attributes_from_excel", path=str(p), key_tag="TAG")
    assert not r["ok"] and r["error"]["code"] == "excel_error" and "TAG" in r["error"]["message"]
    assert call("sync_attributes_from_excel", path=str(tmp_path / "nao.xlsx"), key_tag="TAG")["error"]["code"] == "excel_error"
    assert call("import_blocks_from_excel", path=str(tmp_path / "nao.xlsx"))["error"]["code"] == "excel_error"
    assert call("sync_attributes_from_excel", path=str(p), key_tag=" ")["error"]["code"] == "invalid_parameter"
