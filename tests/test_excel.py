"""Camada Excel pura (sem AutoCAD): leitura, mapeamento, normalização de linhas, caminhos."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from autocad_mcp.config import Config
from autocad_mcp.errors import ExcelIntegrationError, InvalidParameterError
from autocad_mcp.excel_io import build_rows, check_path, read_table, resolve_mapping, write_table


def make_xlsx(path: Path, rows, sheet="Suportes", extra_sheet=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    if extra_sheet:
        wb.create_sheet(extra_sheet).append(["a"])
    wb.save(path)
    return path


def test_read_table_skips_blank_rows_keeps_excel_row_numbers(tmp_path):
    p = make_xlsx(tmp_path / "a.xlsx", [["TAG", "X", "Y"], ["S1", 1, 2], [None, None, None], ["S2", 3, 4]])
    t = read_table(p, None, 1, 100)
    assert t.headers == ["TAG", "X", "Y"] and [r for r, _ in t.rows] == [2, 4]
    assert t.rows[1][1] == {"TAG": "S2", "X": 3, "Y": 4}


def test_read_table_sheet_selection_header_row_and_errors(tmp_path):
    p = make_xlsx(tmp_path / "a.xlsx", [["titulo"], ["TAG", "X"], ["S1", 5]], extra_sheet="Outra")
    assert read_table(p, "suportes", 2, 10).rows[0][1]["X"] == 5
    with pytest.raises(ExcelIntegrationError, match="Aba 'nao'"):
        read_table(p, "nao", 2, 10)
    with pytest.raises(InvalidParameterError):
        read_table(p, None, 0, 10)
    dup = make_xlsx(tmp_path / "d.xlsx", [["TAG", "tag"], ["a", "b"]])
    with pytest.raises(ExcelIntegrationError, match="duplicados"):
        read_table(dup, None, 1, 10)
    big = make_xlsx(tmp_path / "b.xlsx", [["TAG"]] + [[f"S{i}"] for i in range(20)])
    with pytest.raises(ExcelIntegrationError, match="mais de 10"):
        read_table(big, None, 1, 10)
    bad = tmp_path / "corrupt.xlsx"
    bad.write_text("não sou um zip")
    with pytest.raises(ExcelIntegrationError, match="abrir"):
        read_table(bad, None, 1, 10)


def test_auto_mapping_recognises_pt_and_en_names_and_treats_rest_as_attributes():
    m = resolve_mapping(["Bloco", "Camada", "X", "Y", "Z", "Rotação", "TAG", "Tipo Suporte", "ATTR:X"], None)
    assert m["Bloco"] == "block" and m["Camada"] == "layer" and m["X"] == "x" and m["Rotação"] == "rotation"
    assert m["TAG"] == "attr:TAG" and m["Tipo Suporte"] == "attr:TIPO SUPORTE" and m["ATTR:X"] == "attr:X"


def test_explicit_mapping_and_validation():
    m = resolve_mapping(["Suporte", "Este", "Norte"], {"suporte": "attr:tag", "ESTE": "x", "norte": "Y"})
    assert m == {"Suporte": "attr:TAG", "Este": "x", "Norte": "y"}
    with pytest.raises(InvalidParameterError, match="não existe"):
        resolve_mapping(["A"], {"B": "x"})
    with pytest.raises(InvalidParameterError, match="inválido"):
        resolve_mapping(["A"], {"A": "banana"})
    with pytest.raises(InvalidParameterError):
        resolve_mapping(["A"], {"A": "attr:"})


def test_build_rows_numbers_decimal_comma_blanks_and_row_errors(tmp_path):
    p = make_xlsx(tmp_path / "a.xlsx", [
        ["BLOCO", "TAG", "X", "Y", "Z", "ESCALA", "ESCALA_Z", "LINHA"],
        ["SUP", "S1", "1.234,5", "2,5", None, 2, 3, "L1"],
        ["SUP", "S2", "abc", 1, 0, None, None, None],
        [None, "S3", 1, 2, 3, None, None, ""],
    ])
    t = read_table(p, None, 1, 100)
    specs, errs = build_rows(t, resolve_mapping(t.headers, None), default_block="PADRAO")
    assert [s.row for s in specs] == [2, 4] and errs[0]["row"] == 3 and "abc" in errs[0]["message"]
    s1 = specs[0]
    assert (s1.x, s1.y, s1.z) == (1234.5, 2.5, None) and s1.scale == (2.0, 2.0, 3.0) and s1.attributes == {"TAG": "S1", "LINHA": "L1"}
    assert s1.point == (1234.5, 2.5, 0.0) and s1.block == "SUP"
    assert specs[1].block == "PADRAO" and specs[1].attributes == {"TAG": "S3"}  # célula vazia ignorada


def test_write_then_read_roundtrip_and_info_sheet(tmp_path):
    p = tmp_path / "out" / "x.xlsx"
    p.parent.mkdir()
    write_table(p, "Blocos", ["HANDLE", "TAG"], [["1A", "S1"], ["1B", 5.0]], {"unidade": "mm", "linhas": 2})
    t = read_table(p, None, 1, 10)
    assert t.sheet == "Blocos" and t.rows[0][1] == {"HANDLE": "1A", "TAG": "S1"}
    from openpyxl import load_workbook
    wb = load_workbook(p)
    assert wb.sheetnames == ["Blocos", "INFO"] and wb["INFO"]["B1"].value == "mm"


def test_check_path_rules(tmp_path):
    cfg = Config()
    ok = make_xlsx(tmp_path / "a.xlsx", [["A"]])
    assert check_path(str(ok), cfg, write=False) == ok.resolve()
    for bad in ("", "x.csv", "x.xls"):
        with pytest.raises(ExcelIntegrationError):
            check_path(bad, cfg, write=False)
    with pytest.raises(ExcelIntegrationError, match="não encontrado"):
        check_path(str(tmp_path / "nao.xlsx"), cfg, write=False)
    with pytest.raises(ExcelIntegrationError, match="Só é possível gravar"):
        check_path(str(tmp_path / "a.xlsm"), cfg, write=True)
    jail = Config(excel_allowed_dirs=(tmp_path / "ok",))
    (tmp_path / "ok").mkdir()
    with pytest.raises(ExcelIntegrationError, match="fora das pastas"):
        check_path(str(ok), jail, write=False)
    assert check_path(str(tmp_path / "ok" / "n.xlsx"), jail, write=True).parent.name == "ok"
